"""PDF Playfull Advanced tools (offline, document-local operations).
Operations are isolated from the original PDF Toolkit and never upload documents elsewhere.
Fidelity limitations: PDF->PPTX creates image slides; PDF->XLSX extracts detected tables;
PDF->Word uses pdf2docx when available; complex fonts/layout are not guaranteed.
"""
from __future__ import annotations

import io
import json
import math
import os
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape, quoteattr

import fitz

PDF = "application/pdf"
ZIP = "application/zip"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def safe_name(name: str, fallback="hasil") -> str:
    s = re.sub(r"[^\w .()-]", "_", str(name or ""), flags=re.UNICODE).strip(" .")
    return (s[:90] or fallback)


def files_zip(items: list[tuple[str, bytes]]) -> tuple[bytes, str, str]:
    if not items:
        raise ValueError("Tidak ada hasil yang dapat dibuat.")
    if len(items) == 1:
        filename, body = items[0]
        suffix = Path(filename).suffix.lower()
        mime = {".jpg":"image/jpeg", ".jpeg":"image/jpeg", ".png":"image/png", ".pdf":PDF,
                ".txt":"text/plain", ".md":"text/markdown", ".docx":DOCX,
                ".pptx":PPTX, ".xlsx":XLSX}.get(suffix, "application/octet-stream")
        return body, filename, mime
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for filename, body in items:
            z.writestr(safe_name(filename), body)
    return buf.getvalue(), "hasil_banyak_file.zip", ZIP


def open_pdf(data: bytes, password=""):
    pdf = fitz.open(stream=data, filetype="pdf")
    if pdf.needs_pass and not pdf.authenticate(password):
        pdf.close()
        raise ValueError("PDF terkunci. Masukkan password yang benar.")
    return pdf


def save_pdf(pdf):
    return pdf.tobytes(garbage=4, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)


def color(hexcode, default=(0, 0, 0)):
    s = str(hexcode or "").lstrip("#")
    if len(s) == 6 and re.fullmatch("[0-9a-fA-F]{6}", s):
        return tuple(int(s[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return default


def page_which(o, total):
    raw = o.get("pages", "semua")
    if raw is None or str(raw).strip().lower() in ("", "semua", "all"):
        return list(range(total))
    result = []
    for token in re.split(r"[,;\s]+", str(raw)):
        if not token: continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", token)
        if not m: raise ValueError(f"Format halaman salah: {token}")
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        if not (1 <= a <= total and 1 <= b <= total): raise ValueError("Nomor halaman melebihi jumlah halaman PDF.")
        result.extend(range(a - 1, b if a <= b else b - 2, 1 if a <= b else -1))
    return list(dict.fromkeys(result))


def rect_from_obj(op, page):
    vals = [float(op.get(k, 0)) for k in ("x", "y", "w", "h")]
    x, y, w, h = vals
    if not all(math.isfinite(v) for v in vals) or w <= 0 or h <= 0:
        raise ValueError("Koordinat objek tidak valid.")
    r = fitz.Rect(x, y, x + w, y + h) & page.rect
    if r.is_empty or r.width < 1 or r.height < 1:
        raise ValueError("Objek berada di luar halaman.")
    return r


def _is_scanned_page(page) -> bool:
    """Conservatively detect a page mostly covered by a raster scan."""
    area = max(1, page.rect.get_area())
    try:
        return any(fitz.Rect(info["bbox"]).get_area() / area >= .68 for info in page.get_image_info())
    except Exception:
        return False


def _ink_color_near(page, rect, fallback):
    """Estimate ink color from neighboring scan pixels without changing background."""
    try:
        from PIL import Image
        import numpy as np
        margin = max(15, min(65, rect.width * .15))
        sample = (fitz.Rect(rect.x0-margin, rect.y0-margin, rect.x1+margin, rect.y1+margin) & page.rect)
        pix = page.get_pixmap(matrix=fitz.Matrix(.7,.7), clip=sample, colorspace=fitz.csRGB, alpha=False)
        arr = np.asarray(Image.frombytes("RGB", [pix.width,pix.height], pix.samples), dtype=np.float32)
        gray = arr.mean(axis=2)
        baseline = float(np.percentile(gray, 85))
        dark = arr[gray < min(170.0, baseline-35.0)]
        if len(dark) >= 5:
            v=np.median(dark,axis=0)/255.0
            return tuple(float(x) for x in v)
    except Exception:
        pass
    return fallback


def _soft_scan_text(page, rect, value, fontsize, family, bold, italic, ink, align):
    """Slightly soften only NEW text on scan-like pages, leave all other PDF objects untouched.

    Text is embedded as transparent PNG, with an invisible searchable layer if it fits.
    Scan source lettering already burned into the image is NOT magically erased.
    """
    from PIL import Image, ImageDraw, ImageFont, ImageFilter
    scale = 1.75
    w, h = max(2, int(rect.width*scale)), max(2, int(rect.height*scale))
    if w*h > 2_500_000:
        raise ValueError("Kotak teks terlalu besar untuk penyamaan kualitas scan.")
    fontfile = ('DejaVuSerif' if family=='serif' else 'DejaVuSansMono' if family=='mono' else 'DejaVuSans')
    fontfile += ('-BoldOblique' if bold and italic else '-Bold' if bold else '-Oblique' if italic else '') + '.ttf'
    try:
        font = ImageFont.truetype(fontfile, max(8, int(fontsize*scale)))
    except OSError:
        font = ImageFont.load_default()
    overlay = Image.new('RGBA',(w,h),(0,0,0,0))
    draw=ImageDraw.Draw(overlay)
    rgba=tuple(max(0,min(255,round(x*255))) for x in ink)+(220,)
    spacing=max(0,round(fontsize*scale*.2))
    align_map={0:'left',1:'center',2:'right'}
    try:
        draw.multiline_text((0,0), value, font=font, fill=rgba,spacing=spacing,align=align_map.get(align,'left'))
    except ValueError:
        draw.text((0,0),value.splitlines()[0],font=font,fill=rgba)
    overlay=overlay.filter(ImageFilter.GaussianBlur(radius=.42))
    bio=io.BytesIO();overlay.save(bio,format='PNG')
    page.insert_image(rect,stream=bio.getvalue(),overlay=True,keep_proportion=False)
    # Keep inserted wording searchable; invisible text has no sharper visual footprint.
    try:
        face={'sans':'helv','serif':'tiro','mono':'cour'}[family]
        page.insert_textbox(rect,value,fontsize=fontsize,fontname=face,render_mode=3,align=align,overlay=True)
    except Exception:
        pass


def apply_edits(data: bytes, opts: dict) -> bytes:
    """Actual PDF text removal uses redact annotations, never only paints white rectangles."""
    operations = opts.get("operations", [])
    if not isinstance(operations, list) or len(operations) > 600:
        raise ValueError("Maksimum 600 perubahan setiap proses.")
    with open_pdf(data, opts.get("password", "")) as pdf:
        for op in operations:
            kind = op.get("type")
            ix = int(op.get("page", 1)) - 1
            if ix < 0 or ix >= len(pdf): raise ValueError("Halaman perubahan tidak valid.")
            p = pdf[ix]
            rect = rect_from_obj(op, p)
            if kind in ("replace_text", "delete_text", "delete_image", "redact"):
                fill = (0, 0, 0) if kind == "redact" else None
                p.add_redact_annot(rect, fill=fill, cross_out=False)
                # For text-only edits, keep underlying vector images and lines intact.
                # Images=2: blank overlapping pixels only for image deletion/redaction.
                if kind in ("replace_text", "delete_text"):
                    p.apply_redactions(images=0, graphics=0, text=0)
                else:
                    p.apply_redactions(images=2, graphics=1, text=0)
            if kind in ("text", "replace_text"):
                fontsize = min(100, max(5, float(op.get("fontSize", 12))))
                value = str(op.get("text", ""))[:8000]
                family = op.get("font", "sans")
                if family not in ("serif", "mono", "sans"):
                    raise ValueError("Jenis font tidak tersedia.")
                bold = bool(op.get("bold", False))
                italic = bool(op.get("italic", False))
                faces = {
                    "sans": ("helv", "hebo", "heit", "hebi"),
                    "serif": ("tiro", "tibo", "tiit", "tibi"),
                    "mono": ("cour", "cobo", "coit", "cobi"),
                }
                font = faces[family][int(bold) + 2 * int(italic)]
                alignment = {"left": 0, "center": 1, "right": 2}.get(op.get("align"), 0)
                ink=color(op.get("color", "#17324d"))
                # Quality Match applies only to scans and to newly rendered text,
                # not the original page. Native PDFs retain sharp selectable vectors.
                if opts.get("quality_match", False) and _is_scanned_page(p):
                    ink=_ink_color_near(p,rect,ink)
                    _soft_scan_text(p,rect,value,fontsize,family,bold,italic,ink,alignment)
                else:
                    available = p.insert_textbox(rect, value, fontsize=fontsize, fontname=font,
                                                 color=ink, align=alignment, overlay=True)
                    if available < 0:
                        p.insert_text((rect.x0, min(rect.y1, rect.y0 + fontsize * .95)), value.splitlines()[0],
                                      fontsize=fontsize, fontname=font, color=ink, overlay=True)
            elif kind == "move_original_image":
                orig = op.get("original", {})
                old_rect = rect_from_obj(orig, p)
                pix = p.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), clip=old_rect, alpha=False)
                p.add_redact_annot(old_rect, fill=None, cross_out=False)
                p.apply_redactions(images=2, graphics=0, text=0)
                p.insert_image(rect, stream=pix.tobytes("png"), keep_proportion=True, overlay=True)
            elif kind in ("image", "replace_image"):
                import base64
                img = op.get("image", "")
                if not str(img).startswith("data:image/"):
                    raise ValueError("Gambar harus berupa PNG atau JPG.")
                binary = base64.b64decode(str(img).split(",", 1)[-1], validate=True)
                if len(binary) > 5 * 1024 * 1024: raise ValueError("Ukuran gambar maksimum 5 MB.")
                if kind == "replace_image":
                    p.add_redact_annot(rect, fill=None, cross_out=False)
                    p.apply_redactions(images=2, graphics=0, text=0)
                p.insert_image(rect, stream=binary, keep_proportion=True, overlay=True)
            elif kind in ("rectangle", "ellipse"):
                stroke = color(op.get("color", "#2563eb"))
                if kind == "rectangle": p.draw_rect(rect, color=stroke, width=2, overlay=True)
                else: p.draw_oval(rect, color=stroke, width=2, overlay=True)
            elif kind == "highlight":
                shape = p.new_shape()
                shape.draw_rect(rect)
                shape.finish(color=None, fill=color(op.get("color", "#ffff00"), (1,1,0)), fill_opacity=0.25)
                shape.commit(overlay=True)
            elif kind in ("delete_text", "delete_image", "redact"):
                pass
            else:
                raise ValueError(f"Tindakan edit tidak dikenal: {kind}")
        return save_pdf(pdf)


def inspect_page(data: bytes, page_index=1, password=""):
    with open_pdf(data, password) as d:
        i = int(page_index) - 1
        if not (0 <= i < len(d)): raise ValueError("Halaman tidak valid.")
        page = d[i]
        objects = []
        for block in page.get_text("dict").get("blocks", []):
            bbox = block.get("bbox", ())
            if len(bbox) != 4: continue
            x0,y0,x1,y1 = [float(x) for x in bbox]
            if block.get("type") == 0:
                text = "\n".join("".join(span.get("text", "") for span in line.get("spans", [])) for line in block.get("lines", []))
                sizes = [float(span.get("size", 12)) for l in block.get("lines", []) for span in l.get("spans", [])]
                if text.strip():
                    spans = [sp for line in block.get("lines", []) for sp in line.get("spans", [])]
                    lead = spans[0] if spans else {}
                    font_name = str(lead.get("font", "")).lower()
                    family = "mono" if any(t in font_name for t in ("cour", "mono")) else "serif" if any(t in font_name for t in ("times", "serif", "roman")) else "sans"
                    flags = int(lead.get("flags", 0))
                    rgb = int(lead.get("color", 0)) & 0xFFFFFF
                    objects.append({"type":"text","text":text,"x":x0,"y":y0,"w":x1-x0,"h":y1-y0,
                                    "fontSize":round(max(sizes, default=12),1), "font":family,
                                    "bold":bool(flags & 16), "italic":bool(flags & 2),
                                    "color":f"#{rgb:06x}"})
            elif block.get("type") == 1:
                objects.append({"type":"image","x":x0,"y":y0,"w":x1-x0,"h":y1-y0})
        return {"width":page.rect.width,"height":page.rect.height,"pages":len(d),"objects":objects}


def watermark(data: bytes, opts: dict) -> bytes:
    with open_pdf(data, opts.get("password", "")) as d:
        angle = int(opts.get("angle", 0))
        text = str(opts.get("text", "DRAFT"))[:200]
        for i in page_which(opts, len(d)):
            p=d[i]; r=p.rect
            x,y = r.width*0.15,r.height*0.52
            fs = min(72, max(8, int(opts.get("fontSize", 42))))
            sh = p.new_shape(); sh.insert_text((x,y), text, fontsize=fs, rotate=angle if angle in (0,90,180,270) else 0,
                                               color=color(opts.get("color", "#808080")), fill_opacity=0.32)
            sh.commit(overlay=True)
        return save_pdf(d)


def crop(data:bytes, opts: dict) -> bytes:
    with open_pdf(data, opts.get("password", "")) as d:
        for i in page_which(opts, len(d)):
            p = d[i]
            l,t,r,b = [max(0,float(opts.get(k, 0))) for k in ("left","top","right","bottom")]
            rect = fitz.Rect(p.rect.x0+l,p.rect.y0+t,p.rect.x1-r,p.rect.y1-b)
            if rect.is_empty or rect.width < 20 or rect.height < 20:
                raise ValueError("Margin pemotongan terlalu besar.")
            p.set_cropbox(rect)
        return save_pdf(d)


def redact(data: bytes, opts: dict) -> bytes:
    with open_pdf(data, opts.get("password", "")) as d:
        phrase = str(opts.get("phrase", "")).strip()
        areas = opts.get("areas", [])
        count = 0
        for i in page_which(opts, len(d)):
            p=d[i]
            if phrase:
                for rect in p.search_for(phrase):
                    p.add_redact_annot(rect, fill=color(opts.get("color", "#000000")));count+=1
            for op in areas:
                if int(op.get("page",0)) == i+1:
                    p.add_redact_annot(rect_from_obj(op,p),fill=(0,0,0));count+=1
            if count:
                p.apply_redactions(images=2,graphics=1,text=0)
        if not count: raise ValueError("Tidak ada area atau teks yang ditemukan untuk disamarkan.")
        return save_pdf(d),count


def repair(data:bytes,opts:dict) -> bytes:
    try:
        with open_pdf(data,opts.get("password","")) as d:
            return save_pdf(d)
    except Exception as exc:
        raise ValueError("PDF terlalu rusak atau tidak dapat dibuka; pemulihan tidak dijamin.") from exc


def text_to_markdown(data,opts):
    chunks=[]
    with open_pdf(data,opts.get("password","")) as d:
        for j,p in enumerate(d):
            chunks.append(f"<!-- Halaman {j+1} -->")
            blocks = p.get_text("dict").get("blocks",[])
            spans = [s for b in blocks if b.get("type")==0 for line in b.get("lines",[]) for s in line.get("spans",[]) if s.get("text","").strip()]
            avg = sorted([s.get("size",10) for s in spans])[len(spans)//2] if spans else 11
            for b in blocks:
                if b.get("type")!=0:continue
                lines=b.get("lines",[])
                text="\n".join("".join(s.get("text","") for s in line.get("spans",[])) for line in lines).strip()
                if not text:continue
                max_font = max((s.get("size",10) for line in lines for s in line.get("spans",[])),default=10)
                if len(text)<130 and max_font>avg*1.5:chunks.append("# "+text)
                elif len(text)<130 and max_font>avg*1.22:chunks.append("## "+text)
                else:chunks.append(text)
            chunks.append("")
    return "\n\n".join(chunks).encode("utf-8")


def extract_images(data,opts):
    imgs=[]
    with open_pdf(data,opts.get("password","")) as d:
        seen=set()
        for i in page_which(opts,len(d)):
            for j,tup in enumerate(d[i].get_images(full=True),1):
                xref=tup[0]
                if xref in seen:continue
                seen.add(xref)
                detail=d.extract_image(xref)
                if detail and detail.get("image"):
                    imgs.append((f"gambar_halaman_{i+1:03d}_{j:03d}.{detail['ext']}",detail["image"]))
    if not imgs:raise ValueError("Tidak ada gambar bitmap yang dapat diekstrak.")
    return files_zip(imgs)


def libreoffice(data:bytes,src_ext:str,out_ext="pdf",filter_arg=None):
    suffixes={".docx",".doc",".odt",".ppt",".pptx",".odp",".xls",".xlsx",".ods",".html",".htm"}
    if src_ext.lower() not in suffixes: raise ValueError("Ekstensi dokumen tidak didukung.")
    with tempfile.TemporaryDirectory(prefix="pdfplay-lo-") as work:
        src=Path(work)/("input"+src_ext)
        outdir=Path(work)/"output";outdir.mkdir()
        profile=Path(work)/"profile"
        src.write_bytes(data)
        conv=out_ext+(":"+filter_arg if filter_arg else "")
        args=["soffice",f"-env:UserInstallation=file://{profile}","--headless", "--convert-to",conv,
              "--outdir",str(outdir),str(src)]
        try:
            result=subprocess.run(args,capture_output=True,text=True,timeout=90,check=False)
        except (FileNotFoundError,subprocess.TimeoutExpired) as e:
            raise ValueError("LibreOffice tidak tersedia atau proses konversi terlalu lama.") from e
        outputs=list(outdir.glob("*."+out_ext))
        if result.returncode != 0 or not outputs:
            raise ValueError("Konversi LibreOffice gagal: "+(result.stderr or result.stdout)[-350:])
        return outputs[0].read_bytes()


def office_to_pdf(data:bytes, filename:str):
    ext=Path(filename).suffix.lower()
    return libreoffice(data,ext)


def pdf_to_pdfa(data:bytes,opts):
    # LibreOffice Draw imports most normal PDFs and exports PDF/A-2b using SelectPdfVersion=2.
    # Rendering fidelity and PDF/A conformance must be checked with veraPDF externally.
    with tempfile.TemporaryDirectory(prefix="pdfplay-pdfa-") as work:
        src=Path(work)/"input.pdf";src.write_bytes(data)
        out=Path(work)/"out";out.mkdir()
        profile=Path(work)/"profile"
        filter_json=json.dumps({"SelectPdfVersion":{"type":"long","value":"2"}})
        args=["soffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to",
              'pdf:draw_pdf_Export:'+filter_json,"--outdir",str(out),str(src)]
        try: p=subprocess.run(args,capture_output=True,text=True,timeout=90)
        except (FileNotFoundError,subprocess.TimeoutExpired) as e:raise ValueError("LibreOffice tidak tersedia.") from e
        files=list(out.glob("*.pdf"))
        if p.returncode or not files:raise ValueError("Konversi PDF/A tidak berhasil di LibreOffice.")
        return files[0].read_bytes()


def pdf_to_docx(data,opts):
    try:
        from pdf2docx import Converter
    except ImportError:
        raise ValueError("Modul pdf2docx belum terinstal.")
    with tempfile.TemporaryDirectory(prefix="pdfplay-docx-") as w:
        ip=Path(w)/"input.pdf";out=Path(w)/"output.docx";ip.write_bytes(data)
        conv=Converter(str(ip))
        try:conv.convert(str(out),start=0,end=None)
        finally:conv.close()
        if not out.exists():raise ValueError("Konversi PDF ke Word gagal.")
        return out.read_bytes()


def pdf_to_pptx(data,opts):
    from pptx import Presentation
    from pptx.util import Inches
    prs=Presentation()
    prs.slide_width=Inches(10)
    prs.slide_height=Inches(7.5)
    with open_pdf(data,opts.get("password","")) as d:
        if len(d)>80:raise ValueError("PDF → PPTX maksimum 80 halaman.")
        for p in d:
            pix=p.get_pixmap(matrix=fitz.Matrix(1.2,1.2),alpha=False)
            slide=prs.slides.add_slide(prs.slide_layouts[6])
            w,h=prs.slide_width,prs.slide_height
            ratio=min(w/p.rect.width,h/p.rect.height)
            pw=int(p.rect.width*ratio);ph=int(p.rect.height*ratio)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), (w-pw)//2,(h-ph)//2,width=pw,height=ph)
    buf=io.BytesIO();prs.save(buf);return buf.getvalue()


def _xlsx_zip(sheets):
    """Minimal Office Open XML XLSX containing extracted table values (not editable PDF layout)."""
    if not sheets: raise ValueError("Tidak ada tabel yang terdeteksi pada PDF.")
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,"w",zipfile.ZIP_DEFLATED) as z:
        overrides=''.join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' for i in range(1,len(sheets)+1))
        z.writestr('[Content_Types].xml','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'+overrides+'</Types>')
        z.writestr('_rels/.rels','<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr('xl/workbook.xml','<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'+''.join(f'<sheet name={quoteattr(name[:31])} sheetId="{i}" r:id="rId{i}"/>' for i,(name,_) in enumerate(sheets,1))+'</sheets></workbook>')
        z.writestr('xl/_rels/workbook.xml.rels','<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'+''.join(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in range(1,len(sheets)+1))+'</Relationships>')
        for i,(name,table) in enumerate(sheets,1):
            rows=[]
            for r,row in enumerate(table[:10000],1):
                cells=[]
                for col,val in enumerate(row[:80],1):
                    colname="";n=col
                    while n: n,mod=divmod(n-1,26);colname=chr(65+mod)+colname
                    ref=f"{colname}{r}"
                    s=escape(str(val if val is not None else "")[:20000]);cells.append(f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{s}</t></is></c>')
                rows.append(f'<row r="{r}">'+''.join(cells)+'</row>')
            z.writestr(f'xl/worksheets/sheet{i}.xml','<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'+''.join(rows)+'</sheetData></worksheet>')
    return buf.getvalue()


def pdf_to_xlsx(data,opts):
    sheets=[]
    with open_pdf(data,opts.get("password","")) as d:
        for ix,p in enumerate(d):
            try:tables=p.find_tables().tables
            except (AttributeError,ValueError):tables=[]
            for n,t in enumerate(tables):
                matrix=t.extract()
                if matrix:sheets.append((f"Hal{ix+1}_Tabel{n+1}",matrix))
    return _xlsx_zip(sheets[:100])


def ocr_pdf(data,opts):
    """Add a searchable OCR layer to scanned pages without changing already searchable pages."""
    import pytesseract
    from PIL import Image
    lang=str(opts.get("lang","eng"))
    if lang not in ("eng","ind","ind+eng"):raise ValueError("Bahasa OCR tidak didukung.")
    dst=fitz.open()
    with open_pdf(data,opts.get("password","")) as src:
        if len(src)>40:raise ValueError("OCR maksimum 40 halaman per proses pada paket ringan.")
        for i,p in enumerate(src):
            if p.get_text().strip():
                dst.insert_pdf(src,from_page=i,to_page=i);continue
            pix=p.get_pixmap(matrix=fitz.Matrix(1.6,1.6),alpha=False)
            im=Image.open(io.BytesIO(pix.tobytes("png")))
            try:ocr_bytes=pytesseract.image_to_pdf_or_hocr(im,extension="pdf",lang=lang)
            except pytesseract.TesseractError as exc:raise ValueError("Bahasa OCR belum tersedia di server.") from exc
            with fitz.open(stream=ocr_bytes,filetype="pdf") as r:dst.insert_pdf(r)
    try:return save_pdf(dst)
    finally:dst.close()


def add_signature_image(data,opts):
    # Visual signature is not a cryptographic digital signature.
    import base64
    img=opts.get("image")
    if not img or not str(img).startswith("data:image/"):raise ValueError("Gambar tanda tangan harus diunggah.")
    raw=base64.b64decode(str(img).split(",",1)[1],validate=True)
    with open_pdf(data,opts.get("password","")) as d:
        i=int(opts.get("page",1))-1
        if i<0 or i>=len(d):raise ValueError("Halaman tanda tangan tidak ditemukan.")
        r=rect_from_obj(opts,d[i])
        d[i].insert_image(r,stream=raw,keep_proportion=True,overlay=True)
        return save_pdf(d)


def form_field(data,opts):
    with open_pdf(data,opts.get("password","")) as d:
        i=int(opts.get("page",1))-1
        if i<0 or i>=len(d):raise ValueError("Halaman tidak ditemukan.")
        p=d[i];r=rect_from_obj(opts,p)
        widget=fitz.Widget()
        widget.field_name=safe_name(opts.get("field_name","kolom"))
        widget.field_label=widget.field_name
        widget.field_type=fitz.PDF_WIDGET_TYPE_TEXT
        widget.rect=r
        widget.text_font="Helv"
        widget.text_fontsize=min(20,max(8,int(opts.get("fontSize",11))))
        widget.field_value=str(opts.get("default",""))[:200]
        p.add_widget(widget)
        return save_pdf(d)


def compare_pdfs(a,b,opts):
    """Generate a visual side-by-side PDF diff report (no AI)."""
    out=fitz.open()
    with open_pdf(a,opts.get("password","")) as d1,open_pdf(b,opts.get("password2","")) as d2:
        n=max(len(d1),len(d2))
        if n>60:raise ValueError("Bandingkan maksimum 60 halaman.")
        for i in range(n):
            sheet=out.new_page(width=900,height=650)
            sheet.insert_text((26,27),f"Bandingkan PDF — Halaman {i+1}",fontsize=16)
            for j,d in enumerate((d1,d2)):
                if i>=len(d):continue
                p=d[i]
                pix=p.get_pixmap(matrix=fitz.Matrix(0.8,0.8),alpha=False)
                region=fitz.Rect(20+j*450,45,430+j*450,625)
                aspect=min(region.width/pix.width,region.height/pix.height)
                dst=fitz.Rect(region.x0,region.y0,region.x0+pix.width*aspect,region.y0+pix.height*aspect)
                sheet.insert_image(dst,stream=pix.tobytes("jpeg"))
                sheet.insert_text((region.x0,640),"Dokumen A" if j==0 else "Dokumen B",fontsize=10)
    try:return save_pdf(out)
    finally:out.close()


def sign_certificate(data, p12, opts):
    """Cryptographic certificate-based PDF signature, distinct from visible signature images."""
    try:
        from pyhanko.sign import signers
        from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    except ImportError as e:
        raise ValueError("pyHanko belum terinstal di server.") from e
    password = str(opts.get("certificate_password", "")).encode("utf-8")
    with tempfile.TemporaryDirectory(prefix="pdfplay-cert-") as w:
        cert=Path(w)/"signer.p12";cert.write_bytes(p12)
        signer=signers.SimpleSigner.load_pkcs12(pfx_file=str(cert),passphrase=password)
        if signer is None:raise ValueError("Tidak bisa memuat sertifikat PKCS#12 atau password tidak tepat.")
        output=io.BytesIO()
        writer=IncrementalPdfFileWriter(io.BytesIO(data))
        field_name="PDFPlayfullSignature"
        signers.PdfSigner(signers.PdfSignatureMetadata(field_name=field_name),signer=signer).sign_pdf(writer,output=output)
        return output.getvalue()


def detect_form_fields(data,opts):
    """Heuristic: add input boxes after short colon-terminated labels on page."""
    total=0
    with open_pdf(data,opts.get("password","")) as d:
        for p in d:
            existing={w.field_name for w in (p.widgets() or [])}
            lines=[]
            for block in p.get_text("dict").get("blocks",[]):
                if block.get("type") != 0:continue
                for line in block.get("lines",[]):
                    txt="".join(s.get("text","") for s in line.get("spans",[])).strip()
                    if re.match(r"^[^:]{3,50}:\s*$",txt):lines.append((txt,line["bbox"]))
            for txt,bbox in lines[:40]:
                if total>=100:break
                x,y,x1,y1=bbox
                rect=fitz.Rect(x1+5,y, min(p.rect.width-12,x1+210), max(y+22,y1+3))
                if rect.width<65:continue
                w=fitz.Widget();w.field_type=fitz.PDF_WIDGET_TYPE_TEXT;w.rect=rect
                base=safe_name(txt[:-1].lower().replace(" ","_"))
                idx=1;name=base
                while name in existing:name=base+f"_{idx}";idx+=1
                w.field_name=name;w.field_label=txt;w.text_font="Helv";w.text_fontsize=10
                p.add_widget(w);existing.add(name);total+=1
        if total==0:raise ValueError("Tidak ditemukan label pendek berakhiran ':' yang sesuai untuk field otomatis.")
        return save_pdf(d),total


def fill_form_fields(data, opts):
    values=opts.get("fields",{})
    if not isinstance(values,dict):raise ValueError("Data isian tidak valid.")
    found=0
    with open_pdf(data,opts.get("password","")) as d:
        for p in d:
            for w in p.widgets() or []:
                if w.field_name in values:
                    w.field_value=str(values[w.field_name])[:1000];w.update();found+=1
        if not found:raise ValueError("Tidak ada kolom formulir yang cocok.")
        return save_pdf(d),found

def dispatch(tool:str, fs:list, opts:dict):
    if not fs:raise ValueError("Pilih file terlebih dahulu.")
    data=fs[0].data
    base=safe_name(opts.get("nama") or Path(fs[0].name).stem)
    if tool == "edit": result=apply_edits(data,opts);ext=".pdf";mime=PDF
    elif tool == "watermark":result=watermark(data,opts);ext=".pdf";mime=PDF
    elif tool == "crop":result=crop(data,opts);ext=".pdf";mime=PDF
    elif tool == "redact":result,n=redact(data,opts);ext=".pdf";mime=PDF
    elif tool == "repair":result=repair(data,opts);ext=".pdf";mime=PDF
    elif tool == "markdown":result=text_to_markdown(data,opts);ext=".md";mime="text/markdown"
    elif tool == "extract_images":
        body,name,mime=extract_images(data,opts)
        return body,(base+".zip" if name.endswith(".zip") else base+Path(name).suffix),mime,"Gambar diekstrak"
    elif tool == "office_pdf":result=office_to_pdf(data,fs[0].name);ext=".pdf";mime=PDF
    elif tool == "html_pdf":result=libreoffice(data,".html");ext=".pdf";mime=PDF
    elif tool == "pdfa":result=pdf_to_pdfa(data,opts);ext=".pdf";mime=PDF
    elif tool == "pdf_word":result=pdf_to_docx(data,opts);ext=".docx";mime=DOCX
    elif tool == "pdf_pptx":result=pdf_to_pptx(data,opts);ext=".pptx";mime=PPTX
    elif tool == "pdf_xlsx":result=pdf_to_xlsx(data,opts);ext=".xlsx";mime=XLSX
    elif tool == "ocr":result=ocr_pdf(data,opts);ext=".pdf";mime=PDF
    elif tool == "sign_visual":result=add_signature_image(data,opts);ext=".pdf";mime=PDF
    elif tool == "form":result=form_field(data,opts);ext=".pdf";mime=PDF
    elif tool == "form_auto":result,n=detect_form_fields(data,opts);ext=".pdf";mime=PDF
    elif tool == "form_fill":result,n=fill_form_fields(data,opts);ext=".pdf";mime=PDF
    elif tool == "sign_certificate":
        if len(fs)<2:raise ValueError("Butuh PDF dan berkas sertifikat .p12/.pfx.")
        result=sign_certificate(data,fs[1].data,opts);ext=".pdf";mime=PDF
    elif tool == "compare":
        if len(fs)<2:raise ValueError("Bandingkan memerlukan 2 file PDF.")
        result=compare_pdfs(data,fs[1].data,opts);ext=".pdf";mime=PDF
    else:raise ValueError("Alat belum tersedia: "+tool)
    return result,base+ext,mime,"Selesai"


# ---------------------------------------------------------------------------
# Advanced editor batch export. Editing and page extraction are deliberately
# independent: edits are applied to source page coordinates before selection.
# ---------------------------------------------------------------------------

def _unique_result_name(requested, seen):
    """Avoid overwriting entries when different source PDFs share a filename."""
    name = safe_name(requested, "hasil")
    if not name.lower().endswith(".pdf"):
        name += ".pdf"
    if name.lower() not in seen:
        seen.add(name.lower())
        return name
    stem = name[:-4]
    i = 2
    while f"{stem}_{i}.pdf".lower() in seen:
        i += 1
    result = f"{stem}_{i}.pdf"
    seen.add(result.lower())
    return result


def batch_edit(fs, opts):
    """Batch apply visual edits per source file and then export chosen pages.

    groups are explicitly ordered by the browser and have 1-based page indices;
    every group has its own name. No mutable editing state is stored server-side.
    """
    if not (1 <= len(fs) <= 12):
        raise ValueError("Pilih 1 sampai 12 dokumen PDF.")
    if sum(len(f.data) for f in fs) > 55 * 1024 * 1024:
        raise ValueError("Total ukuran dokumen maksimal 55 MB.")
    doc_opts = opts.get("documents", [])
    groups = opts.get("groups", [])
    style = opts.get("style", "combined")
    if style not in ("combined", "range", "page", "document"):
        raise ValueError("Mode hasil tidak valid.")
    if not isinstance(doc_opts, list) or len(doc_opts) != len(fs):
        raise ValueError("Pilihan dokumen tidak cocok dengan unggahan.")
    if not isinstance(groups, list) or not (1 <= len(groups) <= 350):
        raise ValueError("Jumlah kelompok hasil harus 1 sampai 350.")
    if len(json.dumps(opts)) > 9_000_000:
        raise ValueError("Data edit terlalu besar.")

    import workspace_ops as wops
    edited = []
    try:
        for f, spec in zip(fs, doc_opts):
            if not f.name.lower().endswith(".pdf"):
                raise ValueError("Fitur edit banyak dokumen hanya menerima PDF.")
            operations = spec.get("operations", [])
            password = str(spec.get("password", ""))
            if operations:
                edited_bytes = apply_edits(f.data, {"password": password, "operations": operations, "quality_match": bool(opts.get("quality_match", False))})
                pdf = open_pdf(edited_bytes)
            else:
                pdf = open_pdf(f.data, password)
            try:
                wops.apply_document_actions(pdf, spec.get('pageActions', []))
            except Exception:
                pdf.close()
                raise
            edited.append(pdf)

        validated = []
        total_pages = 0
        for i, grp in enumerate(groups):
            idx = grp.get("fileIndex")
            if type(idx) is not int or not 0 <= idx < len(edited):
                raise ValueError(f"Kelompok hasil {i+1} memiliki sumber tidak valid.")
            pages = grp.get("pages", [])
            if not isinstance(pages, list) or not pages or len(pages) > len(edited[idx]):
                raise ValueError(f"Kelompok hasil {i+1} mempunyai daftar halaman tidak valid.")
            if any(type(n) is not int or not 1 <= n <= len(edited[idx]) for n in pages):
                raise ValueError(f"Kelompok hasil {i+1} berisi nomor halaman tidak valid.")
            total_pages += len(pages)
            if total_pages > 800:
                raise ValueError("Maksimum 800 halaman dalam satu proses.")
            validated.append((idx, pages, grp.get("name", "hasil")))

        if style == "combined":
            with fitz.open() as result:
                for idx, pages, _ in validated:
                    source = edited[idx]
                    for n in pages:
                        result.insert_pdf(source, from_page=n - 1, to_page=n - 1)
                data = save_pdf(result)
            name = safe_name(opts.get("combined_name"), "hasil_edit")
            if not name.lower().endswith(".pdf"):
                name += ".pdf"
            return data, name, PDF, len(validated)

        seen = set()
        output = io.BytesIO()
        if len(validated) == 1:
            idx, pages, requested = validated[0]
            with fitz.open() as dest:
                for n in pages:
                    dest.insert_pdf(edited[idx], from_page=n - 1, to_page=n - 1)
                return save_pdf(dest), _unique_result_name(requested, seen), PDF, 1

        with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
            for idx, pages, requested in validated:
                with fitz.open() as dest:
                    for n in pages:
                        dest.insert_pdf(edited[idx], from_page=n - 1, to_page=n - 1)
                    zf.writestr(_unique_result_name(requested, seen), save_pdf(dest))
        zip_name = safe_name(opts.get("zip_name"), "hasil_edit_banyak")
        if not zip_name.lower().endswith(".zip"):
            zip_name += ".zip"
        return output.getvalue(), zip_name, ZIP, len(validated)
    finally:
        for pdf in edited:
            pdf.close()
